# -*- coding: utf-8 -*-
"""
verify_phase5_4.py
------------------
Automated Verification & Data Quality Audit Suite untuk Fase 5.4:
Instagram Preprocessing & Final Dataset ekosistem Zahir Social Analytics.

Cakupan Pengujian:
- Test 1: Canonical Raw Source Immutability (MD5 verification)
- Test 2: Play Store & Gold Set Immutability (Frozen reference preservation)
- Test 3: Row Completeness & 1-to-1 Preservation (55 rows raw == 55 rows final)
- Test 4: Uniqueness of Post IDs and URLs (0 duplicates)
- Test 5: Engagement Metrics Integrity (Likes & comments match raw byte-for-byte)
- Test 6: Views Column Missingness Integrity (Strictly 100% NaN, not zeroed)
- Test 7: Total Interactions Arithmetic Consistency (total_interactions == likes + comments)
- Test 8: Datetime & Timezone Parsing Validity (UTC to WIB conversion, ISO-8601, calendar fields)
- Test 9: Caption & Text Processing Integrity (Conservative cleaning, whitespace, word counts)
- Test 10: Hashtags Consistency (hashtag_count == 0 when NaN, > 0 when present)
- Test 11: Semantic & Structural Equivalence between CSV and XLSX
- Test 12: Privacy & Identifier Minimization (No commenter PII or private credentials)
- Test 13: Data Quality Audit Artifacts Verification (CSV and 5-sheet XLSX audit reports)
"""

import os
import sys
import hashlib
import pandas as pd
import numpy as np
import openpyxl

print("=" * 75)
print("PHASE 5.4 AUTOMATED VERIFICATION & QUALITY AUDIT SUITE")
print("=" * 75)

errors = []
passes = []

# Paths
RAW_CANONICAL_CSV = os.path.join("data", "raw", "instagram_posts_raw.csv")
RAW_BASELINE_CSV = os.path.join("data", "raw", "instagram_posts_raw_phase1_6_baseline.csv")
FINAL_CSV = os.path.join("data", "final", "instagram_posts_final.csv")
FINAL_XLSX = os.path.join("data", "final", "instagram_posts_final.xlsx")
AUDIT_CSV = os.path.join("data", "final", "instagram_analysis", "instagram_data_quality_audit.csv")
AUDIT_XLSX = os.path.join("data", "final", "instagram_analysis", "instagram_data_quality_audit.xlsx")
GOLD_SET_CSV = os.path.join("data", "final", "playstore_gold_set_final.csv")
PLAYSTORE_FINAL_CSV = os.path.join("data", "final", "playstore_reviews_final.csv")

EXPECTED_RAW_MD5 = "c9e7882eff8282bb698825a0ef712817"
EXPECTED_ROW_COUNT = 55


def compute_md5(file_path: str) -> str:
    hasher = hashlib.md5()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


# ----------------------------------------------------
# Test 1: Canonical Raw Source Immutability
# ----------------------------------------------------
try:
    assert os.path.exists(RAW_CANONICAL_CSV), f"Missing {RAW_CANONICAL_CSV}"
    assert os.path.exists(RAW_BASELINE_CSV), f"Missing {RAW_BASELINE_CSV}"
    
    md5_raw = compute_md5(RAW_CANONICAL_CSV)
    md5_base = compute_md5(RAW_BASELINE_CSV)
    
    assert md5_raw == EXPECTED_RAW_MD5, f"Raw MD5 {md5_raw} != {EXPECTED_RAW_MD5}"
    assert md5_base == EXPECTED_RAW_MD5, f"Baseline MD5 {md5_base} != {EXPECTED_RAW_MD5}"
    assert md5_raw == md5_base, "Raw and baseline baseline hashes do not match"
    
    p = f"[PASS] Test 1: Canonical raw Instagram data is 100% frozen & immutable (MD5: {md5_raw})."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 1: Raw immutability failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 2: Play Store & Gold Set Immutability
# ----------------------------------------------------
try:
    assert os.path.exists(GOLD_SET_CSV), f"Missing {GOLD_SET_CSV}"
    assert os.path.exists(PLAYSTORE_FINAL_CSV), f"Missing {PLAYSTORE_FINAL_CSV}"
    
    df_gold = pd.read_csv(GOLD_SET_CSV)
    df_ps = pd.read_csv(PLAYSTORE_FINAL_CSV)
    
    assert len(df_gold) == 417, f"Gold set length {len(df_gold)} != 417"
    assert len(df_ps) == 417, f"Play Store length {len(df_ps)} != 417"
    
    sent_counts = df_gold["human_sentiment_label"].value_counts().to_dict()
    assert sent_counts.get("positive", 0) == 232, f"Positive count mismatch: {sent_counts}"
    assert sent_counts.get("neutral", 0) == 112, f"Neutral count mismatch: {sent_counts}"
    assert sent_counts.get("negative", 0) == 73, f"Negative count mismatch: {sent_counts}"
    
    p = "[PASS] Test 2: Play Store reviews & Gold Set remain strictly unaltered (417 reviews: 232 pos, 112 neu, 73 neg)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 2: Play Store immutability failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 3: Row Completeness & 1-to-1 Preservation
# ----------------------------------------------------
try:
    assert os.path.exists(FINAL_CSV), f"Missing {FINAL_CSV}"
    df_raw = pd.read_csv(RAW_CANONICAL_CSV)
    df_final = pd.read_csv(FINAL_CSV)
    
    assert len(df_raw) == EXPECTED_ROW_COUNT, f"Raw rows {len(df_raw)} != {EXPECTED_ROW_COUNT}"
    assert len(df_final) == EXPECTED_ROW_COUNT, f"Final rows {len(df_final)} != {EXPECTED_ROW_COUNT}"
    assert len(df_final) == len(df_raw), "Row count discrepancy between raw and final"
    
    # Check that all post_ids from raw are present in final
    assert set(df_raw["post_id"]) == set(df_final["post_id"]), "Set of post_ids does not match 1-to-1"
    
    p = f"[PASS] Test 3: Row completeness confirmed (100% of {EXPECTED_ROW_COUNT} raw posts preserved, 0 dropped)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 3: Row completeness failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 4: Uniqueness of Post IDs and URLs
# ----------------------------------------------------
try:
    assert df_final["post_id"].nunique() == EXPECTED_ROW_COUNT, "Duplicate post_id detected in final dataset"
    assert df_final["post_url"].nunique() == EXPECTED_ROW_COUNT, "Duplicate post_url detected in final dataset"
    assert df_final["post_id"].isna().sum() == 0, "Null post_id detected"
    assert df_final["post_url"].isna().sum() == 0, "Null post_url detected"
    
    p = f"[PASS] Test 4: Uniqueness verified (55 unique post_id, 55 unique post_url, 0 duplicates)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 4: Uniqueness test failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 5: Engagement Metrics Integrity (Likes & Comments)
# ----------------------------------------------------
try:
    raw_likes_map = df_raw.set_index("post_id")["likes"].to_dict()
    final_likes_map = df_final.set_index("post_id")["likes"].to_dict()
    assert raw_likes_map == final_likes_map, "Likes mapping mismatch between raw and final"
    
    raw_comments_map = df_raw.set_index("post_id")["comments"].to_dict()
    final_comments_map = df_final.set_index("post_id")["comments"].to_dict()
    assert raw_comments_map == final_comments_map, "Comments mapping mismatch between raw and final"
    
    sum_likes = int(df_final["likes"].sum())
    sum_comments = int(df_final["comments"].sum())
    assert sum_likes == 12643, f"Sum of likes {sum_likes} != 12643"
    assert sum_comments == 504, f"Sum of comments {sum_comments} != 504"
    assert (df_final["likes"] >= 0).all(), "Negative likes detected"
    assert (df_final["comments"] >= 0).all(), "Negative comments detected"
    
    p = f"[PASS] Test 5: Engagement metrics verified (Total likes: {sum_likes}, Total comments: {sum_comments}; match raw 100%)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 5: Engagement metrics integrity failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 6: Views Column Missingness Integrity (Strictly NaN)
# ----------------------------------------------------
try:
    # In final CSV, views must be 100% NaN (missing)
    assert df_final["views"].isna().all(), "Views column is not 100% NaN in CSV"
    assert len(df_final["views"].dropna()) == 0, "Non-null views values found"
    
    # In final XLSX, views must also be 100% NaN / empty
    df_xlsx_main = pd.read_excel(FINAL_XLSX, sheet_name="Instagram Posts Final")
    assert df_xlsx_main["views"].isna().all(), "Views column is not 100% NaN in XLSX"
    
    p = "[PASS] Test 6: Views metric integrity confirmed (Strictly 100% NaN; platform web unauthenticated truth preserved, not zeroed)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 6: Views missingness failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 7: Total Interactions Arithmetic Consistency
# ----------------------------------------------------
try:
    expected_interactions = df_final["likes"] + df_final["comments"]
    assert (df_final["total_interactions"] == expected_interactions).all(), "total_interactions != likes + comments"
    sum_interactions = int(df_final["total_interactions"].sum())
    assert sum_interactions == 13147, f"Sum of total_interactions {sum_interactions} != 13147"
    
    p = f"[PASS] Test 7: Total interactions arithmetic verified (Sum = {sum_interactions} == 12,643 likes + 504 comments)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 7: Total interactions arithmetic failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 8: Datetime & Timezone Parsing Validity
# ----------------------------------------------------
try:
    dt_utc = pd.to_datetime(df_final["post_datetime_utc"], utc=True)
    dt_wib = pd.to_datetime(df_final["post_datetime_wib"])
    
    # Check that WIB is exactly 7 hours ahead of UTC
    time_diff_hours = (dt_wib.dt.tz_localize(None) - dt_utc.dt.tz_localize(None)).dt.total_seconds() / 3600.0
    assert (time_diff_hours == 7.0).all(), "WIB timezone offset is not exactly +7 hours from UTC"
    
    # Check post_date format YYYY-MM-DD
    assert df_final["post_date"].str.match(r"^\d{4}-\d{2}-\d{2}$").all(), "post_date does not match YYYY-MM-DD format"
    assert (df_final["post_year"] >= 2024).all() and (df_final["post_year"] <= 2026).all(), "post_year out of expected range"
    assert (df_final["post_month"] >= 1).all() and (df_final["post_month"] <= 12).all(), "post_month out of range"
    assert (df_final["post_day"] >= 1).all() and (df_final["post_day"] <= 31).all(), "post_day out of range"
    assert (df_final["post_hour_wib"] >= 0).all() and (df_final["post_hour_wib"] <= 23).all(), "post_hour_wib out of range"
    
    days_of_week = {"Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"}
    assert set(df_final["post_day_of_week"]).issubset(days_of_week), f"Invalid day of week: {set(df_final['post_day_of_week'])}"
    
    p = "[PASS] Test 8: Datetime parsing & WIB (+07:00) standardizations verified across all 55 posts."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 8: Datetime parsing failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 9: Caption & Text Processing Integrity
# ----------------------------------------------------
try:
    # Original raw caption must be preserved 1-to-1
    assert (df_final["caption"] == df_raw["caption"]).all(), "Raw caption text altered"
    assert (df_final["caption_length_raw"] == df_raw["caption_length"]).all(), "caption_length_raw mismatch"
    assert (df_final["caption_length_raw"] >= 115).all(), "Short/empty caption found"
    
    # Cleaned caption must be non-empty and stripped of excess whitespace
    assert (df_final["caption_clean"].str.strip() != "").all(), "Empty caption_clean found"
    assert (df_final["caption_clean"].str.contains(r"\s{2,}") == False).all(), "Double whitespace found in caption_clean"
    assert (df_final["caption_word_count"] > 0).all(), "Zero word count in caption_clean"
    
    p = "[PASS] Test 9: Caption & text processing verified (Original caption preserved 100%, conservative cleaning valid)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 9: Caption processing failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 10: Hashtags Consistency
# ----------------------------------------------------
try:
    # In 16 rows, hashtags is NaN and hashtag_count is 0
    nan_ht = df_final[df_final["hashtags"].isna()]
    assert len(nan_ht) == 16, f"Expected 16 NaN hashtags, got {len(nan_ht)}"
    assert (nan_ht["hashtag_count"] == 0).all(), "NaN hashtags with non-zero hashtag_count"
    assert (nan_ht["has_hashtags"] == False).all(), "NaN hashtags with has_hashtags=True"
    
    # In 39 rows, hashtags is not NaN and hashtag_count > 0
    non_nan_ht = df_final[df_final["hashtags"].notna()]
    assert len(non_nan_ht) == 39, f"Expected 39 non-null hashtags, got {len(non_nan_ht)}"
    assert (non_nan_ht["hashtag_count"] > 0).all(), "Non-null hashtags with zero hashtag_count"
    assert (non_nan_ht["has_hashtags"] == True).all(), "Non-null hashtags with has_hashtags=False"
    
    p = "[PASS] Test 10: Hashtags logic verified (39 posts with hashtags [count 1-16], 16 posts without hashtags)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 10: Hashtags logic failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 11: Semantic & Structural Equivalence between CSV and XLSX
# ----------------------------------------------------
try:
    assert os.path.exists(FINAL_XLSX), f"Missing {FINAL_XLSX}"
    wb = openpyxl.load_workbook(FINAL_XLSX)
    assert "Instagram Posts Final" in wb.sheetnames, "Missing Instagram Posts Final sheet"
    assert "Data Dictionary & Quality" in wb.sheetnames, "Missing Data Dictionary sheet"
    
    df_xlsx_main = pd.read_excel(FINAL_XLSX, sheet_name="Instagram Posts Final")
    assert len(df_final) == len(df_xlsx_main), "Row count mismatch between CSV and XLSX"
    assert list(df_final.columns) == list(df_xlsx_main.columns), "Column list mismatch between CSV and XLSX"
    assert len(df_final.columns) == 27, f"Expected 27 columns, got {len(df_final.columns)}"
    
    # Check numeric columns equivalence
    for col in ["likes", "comments", "total_interactions", "post_year", "post_month", "post_day", "post_hour_wib", "caption_word_count"]:
        assert (df_final[col].values == df_xlsx_main[col].values).all(), f"Value mismatch in {col} between CSV and XLSX"
        
    p = "[PASS] Test 11: CSV and XLSX semantic & structural equivalence confirmed (27 columns, 55 rows, dual sheets)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 11: CSV/XLSX equivalence failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 12: Privacy & Identifier Minimization
# ----------------------------------------------------
try:
    # Verify no commenter usernames, private accounts, or personal user avatars
    prohibited_cols = ["user_name", "user_image", "commenter_id", "commenter_name", "password", "token", "cookie"]
    for c in df_final.columns:
        assert c.lower() not in prohibited_cols, f"Prohibited column {c} found in final dataset"
        
    # Check original_account is restricted to known corporate / creator partner entities
    allowed_accounts = {"zahiraccounting", "muhamad.is", "pecahtelur.id", "zahironlineid", "modalngobrol", "zahirerp"}
    assert set(df_final["original_account"]).issubset(allowed_accounts), f"Unexpected account: {set(df_final['original_account'])}"
    
    p = "[PASS] Test 12: Privacy verified (Zero personal commenter identifiers or credentials; verified brand accounts only)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 12: Privacy test failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Test 13: Data Quality Audit Artifacts Verification
# ----------------------------------------------------
try:
    assert os.path.exists(AUDIT_CSV), f"Missing {AUDIT_CSV}"
    assert os.path.exists(AUDIT_XLSX), f"Missing {AUDIT_XLSX}"
    
    df_audit_csv = pd.read_csv(AUDIT_CSV)
    assert len(df_audit_csv) >= 20, f"Audit summary rows {len(df_audit_csv)} < 20"
    
    wb_audit = openpyxl.load_workbook(AUDIT_XLSX)
    expected_sheets = ["Audit Summary", "Column Completeness", "Content Types & Engagement", "Account & Collaboration", "Engagement Statistics"]
    for s in expected_sheets:
        assert s in wb_audit.sheetnames, f"Missing sheet {s} in audit XLSX"
        
    p = f"[PASS] Test 13: Data quality audit reports verified (CSV: {len(df_audit_csv)} rows, XLSX: 5 comprehensive sheets)."
    passes.append(p)
    print(p)
except Exception as e:
    err = f"[FAIL] Test 13: Audit artifacts failed: {e}"
    errors.append(err)
    print(err)

# ----------------------------------------------------
# Final Summary
# ----------------------------------------------------
print("=" * 75)
print(f"VERIFICATION SUMMARY: {len(passes)} PASSED, {len(errors)} FAILED")
print("=" * 75)

if errors:
    print("\nFAILURES ENCOUNTERED:")
    for err in errors:
        print("  -", err)
    sys.exit(1)
else:
    print("\nALL 13 PHASE 5.4 AUTOMATED INTEGRITY TESTS PASSED SUCCESSFULLY!")
    sys.exit(0)
